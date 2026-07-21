"""End-to-end acceptance tests for the downloadable application workbook.

These tests use local indexed-search fixtures and open the resulting bytes with
openpyxl. They therefore check the joined application path without a live key,
network request or browser automation.
"""

from io import BytesIO

from openpyxl import load_workbook

from election_extractor.extraction import build_extraction_query
from election_extractor.models import (
    DiscoveredElectionArea,
    DiscoveryReport,
    DiscoveryStatus,
    SearchResult,
)
from election_extractor.search_providers.mock_provider import MockSearchProvider
from election_extractor.workflow import run_extraction_workflow


INDEX_URL = (
    "https://mycouncil.surreycc.gov.uk/"
    "mgElectionElectionAreaResults.aspx?EID=16"
)
FIRST_URL = (
    "https://mycouncil.surreycc.gov.uk/"
    "mgElectionAreaResults.aspx?ID=201&RPID=0"
)
SECOND_URL = (
    "https://mycouncil.surreycc.gov.uk/"
    "mgElectionAreaResults.aspx?ID=202&RPID=0"
)


def _area(name: str, url: str) -> DiscoveredElectionArea:
    """Create one discovered area with trusted index metadata."""

    return DiscoveredElectionArea(
        election_year=2021,
        election_name="2021 Surrey County Council election",
        division_ward_name=name,
        result_url=url,
        source_index_url=INDEX_URL,
        discovery_status=DiscoveryStatus.DISCOVERED,
    )


def _complete_result(area: DiscoveredElectionArea) -> SearchResult:
    """Return a small internally consistent published result fixture."""

    return SearchResult(
        title=f"{area.division_ward_name} election result",
        url=area.result_url,
        snippet=" | ".join(
            (
                "Election name: 2021 Surrey County Council election",
                "Election date: 6 May 2021",
                "Authority: Surrey County Council",
                f"Division/Ward: {area.division_ward_name}",
                "Seats: 1",
                "Candidate: Example Candidate",
                "Party: Example Party",
                "Votes: 100",
                "Vote share: 100%",
                "Outcome: Elected",
                "Total votes: 100",
                "Valid votes: 100",
                "Electorate: 200",
                "Ballot papers issued: 100",
                "Ballot papers rejected: 0",
                "Turnout: 50%",
            )
        ),
    )


def _worksheet_summary(sheet) -> dict[str, object]:
    """Read the fixed six-row Voting Summary table from an area worksheet."""

    detail_row = next(
        row
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row, 1).value == "Detail"
    )
    return {
        sheet.cell(row, 1).value: sheet.cell(row, 2).value
        for row in range(detail_row + 1, detail_row + 7)
    }


def test_index_url_produces_every_area_sheet_and_required_audit(monkeypatch) -> None:
    """Acceptance scenario A: one index expands into all discovered areas."""

    areas = (_area("Addlestone", FIRST_URL), _area("Banstead", SECOND_URL))
    provider = MockSearchProvider(
        {
            build_extraction_query(area): (_complete_result(area),)
            for area in areas
        }
    )

    # Discovery itself has separate query-level tests. Here a fixed discovery
    # report isolates the end-to-end extraction, validation and export contract.
    monkeypatch.setattr(
        "election_extractor.workflow.discover_election_areas",
        lambda url, search_provider, indexed_search_only, require_indexed_search: (
            DiscoveryReport(INDEX_URL, areas, ())
        ),
    )

    result = run_extraction_workflow(
        INDEX_URL,
        provider=provider,
        run_targeted_searches=False,
    )
    workbook = load_workbook(BytesIO(result.workbook_bytes))

    assert workbook.sheetnames == ["Index", "Addlestone", "Banstead", "Extraction Log"]
    assert workbook["Index"].max_row == 3
    assert workbook["Index"]["F2"].value == "Complete"
    assert workbook["Index"]["F3"].value == "Complete"
    assert workbook["Index"]["B2"].hyperlink.target == "#'Addlestone'!A1"
    assert workbook["Index"]["E2"].hyperlink.target == FIRST_URL
    assert _worksheet_summary(workbook["Addlestone"])["Total votes"] == 100
    assert result.complete == 2
    assert result.incomplete == 0
    assert result.failed == 0


def test_direct_url_has_exactly_index_area_and_log_sheets() -> None:
    """Acceptance scenario B: a direct URL exports exactly one area."""

    # Direct processing starts without area metadata; the indexed evidence
    # supplies the published ward name used for the worksheet.
    direct_area = DiscoveredElectionArea(
        election_year=None,
        election_name=None,
        division_ward_name=None,
        result_url=FIRST_URL,
        source_index_url=FIRST_URL,
        discovery_status=DiscoveryStatus.MISSING_AREA_NAME,
    )
    published_area = _area("Addlestone", FIRST_URL)
    provider = MockSearchProvider(
        {build_extraction_query(direct_area): (_complete_result(published_area),)}
    )

    result = run_extraction_workflow(
        FIRST_URL,
        provider=provider,
        run_targeted_searches=False,
    )
    workbook = load_workbook(BytesIO(result.workbook_bytes))

    assert workbook.sheetnames == ["Index", "Addlestone", "Extraction Log"]
    assert _worksheet_summary(workbook["Addlestone"]) == {
        "Seats": 1,
        "Total votes": 100,
        "Electorate": 200,
        "Ballot papers issued": 100,
        "Ballot papers rejected": 0,
        "Turnout": 50,
    }


def test_missing_indexed_values_remain_blank_and_are_audited() -> None:
    """Acceptance scenario C: incomplete snippets never create substitute values."""

    direct_area = DiscoveredElectionArea(
        election_year=None,
        election_name=None,
        division_ward_name=None,
        result_url=FIRST_URL,
        source_index_url=FIRST_URL,
        discovery_status=DiscoveryStatus.MISSING_AREA_NAME,
    )
    incomplete = SearchResult(
        title="Addlestone election result",
        url=FIRST_URL,
        snippet=(
            "Election date: 6 May 2021 | Division/Ward: Addlestone | Seats: 1 | "
            "Candidate: Example Candidate | Party: Example Party | Outcome: Elected"
        ),
    )
    provider = MockSearchProvider(
        {build_extraction_query(direct_area): (incomplete,)}
    )

    result = run_extraction_workflow(
        FIRST_URL,
        provider=provider,
        run_targeted_searches=False,
    )
    workbook = load_workbook(BytesIO(result.workbook_bytes))
    index = workbook["Index"]
    area = workbook["Addlestone"]
    candidate_header = next(
        row
        for row in range(1, area.max_row + 1)
        if area.cell(row, 1).value == "Candidate"
    )

    assert index["F2"].value == "Incomplete"
    assert "votes_received" in index["G2"].value
    assert area.cell(candidate_header + 1, 3).value is None
    assert _worksheet_summary(area)["Total votes"] is None
    assert "not inferred or invented" in area["A11"].value.casefold()
    assert all(area.cell(12, column).value is None for column in range(1, 6))
    assert workbook["Extraction Log"].max_row == 2
