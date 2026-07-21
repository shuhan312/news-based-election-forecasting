"""Tests for the lawful archived-copy fallback used when live pages are blocked.

Every test runs against deterministic in-memory doubles: no live council
request, no live Wayback Machine request and no SerpAPI key. The doubles model
the exact behaviour observed in production — Incapsula returns HTTP 200 with a
small challenge stub, while the Wayback CDX index lists HTTP-200 captures of
the same official URLs.
"""

from __future__ import annotations

import json
from io import BytesIO
from urllib.parse import parse_qs, urlsplit

import pytest
from openpyxl import load_workbook

from election_extractor.discovery import OfficialArchiveResponse
from election_extractor.official_archive_fallback import (
    ARCHIVED_OFFICIAL_COPY,
    ArchiveFallbackOfficialArchiveClient,
    ArchiveFallbackOfficialPageClient,
    ArchivedSnapshotNotFoundError,
    WAYBACK_CDX_ENDPOINT,
    WaybackOfficialArchiveClient,
    WaybackOfficialPageClient,
    WaybackSnapshotResolver,
)
from election_extractor.official_source import (
    OfficialPageClassification,
    OfficialPageResponse,
    fetch_and_diagnose_official_page,
)
from election_extractor.search_providers.mock_provider import MockSearchProvider
from election_extractor.workflow import run_extraction_workflow


AREA_URL = (
    "https://mycouncil.surreycc.gov.uk/mgElectionAreaResults.aspx?ID=201&RPID=0"
)
AREA_URL_WITHOUT_RPID = (
    "https://mycouncil.surreycc.gov.uk/mgElectionAreaResults.aspx?ID=201"
)
# EID=99 is deliberately absent from config/elections.json so the end-to-end
# test exercises pure discovery without the configured 81-area denominator.
UNCONFIGURED_INDEX_URL = (
    "https://mycouncil.surreycc.gov.uk/mgElectionElectionAreaResults.aspx?EID=99"
)

INCAPSULA_BODY = (
    "<html><head><title>Request unsuccessful</title></head>"
    "<body>Request unsuccessful. Incapsula incident ID: 50-123</body></html>"
)


def official_page_body(ward_name: str) -> str:
    """Return one complete official result page for ``ward_name``.

    The figures are internally consistent so validation passes: candidate
    votes equal total votes, ballots issued minus rejected equals total votes,
    and turnout matches ballots issued over electorate at published rounding.
    """
    return f"""
    <html><head><title>Election results for {ward_name}, 6 May 2021</title></head>
    <body><h1>County Council Election 2021</h1><p>Surrey County Council</p>
      <table>
        <tr><th>Election Candidate</th><th>Party</th><th>Votes</th><th>Vote Share</th><th>Outcome</th></tr>
        <tr><td>John Raymond Furey</td><td>Conservative</td><td>1,146</td><td>100%</td><td>Elected</td></tr>
      </table>
      <table summary="Voting summary table"><caption>Voting Summary</caption>
        <tr><th>Details</th><th>Number</th></tr>
        <tr><td>Seats</td><td>1</td></tr>
        <tr><td>Total votes</td><td>1,146</td></tr>
        <tr><td>Electorate</td><td>11,109</td></tr>
        <tr><td>Number of ballot papers issued</td><td>1,155</td></tr>
        <tr><td>Number of ballot papers rejected</td><td>9</td></tr>
        <tr><td>Turnout</td><td>10.4%</td></tr>
      </table>
    </body></html>
    """


def snapshot_url(timestamp: str, original_url: str) -> str:
    return f"https://web.archive.org/web/{timestamp}id_/{original_url}"


class FakeWaybackHTTP:
    """Serve CDX rows, snapshot bodies and replay redirects from fixtures."""

    def __init__(
        self,
        cdx: dict[str, list[list[str]]],
        pages: dict[str, tuple[int, str]],
        replay: dict[str, tuple[str, int, str]] | None = None,
    ) -> None:
        # ``cdx`` maps a requested CDX ``url=`` key to data rows (no header);
        # ``pages`` maps a full snapshot URL to ``(status, body)``;
        # ``replay`` maps a fast-path probe URL to the redirect outcome
        # ``(final_url, status, body)`` exactly as urllib reports it after
        # following the Wayback 302 to the newest capture.
        self.cdx = cdx
        self.pages = pages
        self.replay = replay or {}
        self.requests: list[str] = []

    def get(self, url: str) -> tuple[int, str, str]:
        self.requests.append(url)
        if url.startswith(WAYBACK_CDX_ENDPOINT):
            key = parse_qs(urlsplit(url).query)["url"][0]
            rows = self.cdx.get(key)
            if not rows:
                return 200, url, ""
            return 200, url, json.dumps([["timestamp", "original"], *rows])
        if url in self.replay:
            final_url, status, body = self.replay[url]
            return status, final_url, body
        status, body = self.pages.get(url, (404, "Snapshot not found"))
        return status, url, body


class CountingLiveClient:
    """Return a fixed live response while counting how often it is contacted."""

    def __init__(self, response: OfficialPageResponse) -> None:
        self.response = response
        self.calls = 0

    def fetch(self, url: str) -> OfficialPageResponse:
        self.calls += 1
        return self.response


class CountingLiveArchiveClient:
    """Archive-page variant of the counting live client."""

    def __init__(self, response: OfficialArchiveResponse) -> None:
        self.response = response
        self.calls = 0

    def fetch(self, url: str) -> OfficialArchiveResponse:
        self.calls += 1
        return self.response


def protected_live_page_client() -> CountingLiveClient:
    return CountingLiveClient(OfficialPageResponse(200, AREA_URL, INCAPSULA_BODY))


def test_resolver_uses_the_exact_url_key_first_and_decodes_entities() -> None:
    """The first equivalent key with captures supplies the newest-first list."""

    older = "20200101000000"
    newer = "20240101000000"
    # The CDX index stores some originals with HTML-encoded ampersands; the
    # resolver must decode them before constructing replay URLs.
    encoded_original = (
        "https://mycouncil.surreycc.gov.uk/mgElectionAreaResults.aspx"
        "?ID=201&amp;RPID=0"
    )
    http = FakeWaybackHTTP(
        cdx={
            AREA_URL: [[older, encoded_original], [newer, AREA_URL]],
            # This equivalent key must not even be queried once the exact URL
            # has produced captures.
            AREA_URL_WITHOUT_RPID: [["20250101000000", AREA_URL_WITHOUT_RPID]],
        },
        pages={},
    )

    captures = WaybackSnapshotResolver(http).captures(AREA_URL)

    assert [capture.timestamp for capture in captures] == [newer, older]
    # HTML entities decoded: the audit trail must cite the real official URL.
    assert captures[1].original_url == AREA_URL
    assert captures[0].iso_timestamp == "2024-01-01T00:00:00Z"
    assert len(http.requests) == 1


def test_resolver_falls_back_to_rpid_variants_and_caches_the_result() -> None:
    """RPID is a display parameter; captures under any variant are equivalent."""

    timestamp = "20240101000000"
    http = FakeWaybackHTTP(
        cdx={AREA_URL_WITHOUT_RPID: [[timestamp, AREA_URL_WITHOUT_RPID]]},
        pages={},
    )
    resolver = WaybackSnapshotResolver(http)

    first = resolver.captures(AREA_URL)
    requests_after_first = len(http.requests)
    second = resolver.captures(AREA_URL)

    assert [capture.timestamp for capture in first] == [timestamp]
    assert first[0].snapshot_url == snapshot_url(timestamp, AREA_URL_WITHOUT_RPID)
    # The second resolution is answered from the cache without new requests.
    assert second == first
    assert len(http.requests) == requests_after_first


def test_fast_path_serves_the_newest_capture_in_one_request() -> None:
    """The replay redirect resolves and serves a page without any CDX query."""

    real_timestamp = "20210228064005"
    probe = snapshot_url("20991231235959", AREA_URL)
    http = FakeWaybackHTTP(
        cdx={},
        pages={},
        replay={
            probe: (
                snapshot_url(real_timestamp, AREA_URL),
                200,
                official_page_body("Addlestone"),
            )
        },
    )

    response = WaybackOfficialPageClient(http).fetch(AREA_URL)

    assert response.archive_snapshot_timestamp == "2021-02-28T06:40:05Z"
    assert response.archive_snapshot_url == snapshot_url(real_timestamp, AREA_URL)
    assert response.final_url == AREA_URL
    # Exactly one network request: the probe that followed the redirect.
    assert http.requests == [probe]


def test_wayback_page_client_serves_archived_copy_with_full_provenance() -> None:
    timestamp = "20220925060643"
    probe = snapshot_url("20991231235959", AREA_URL)
    http = FakeWaybackHTTP(
        cdx={},
        pages={},
        replay={
            probe: (
                snapshot_url(timestamp, AREA_URL),
                200,
                official_page_body("Addlestone"),
            )
        },
    )

    response = WaybackOfficialPageClient(http).fetch(AREA_URL)

    assert response.retrieval_source == ARCHIVED_OFFICIAL_COPY
    assert response.archive_snapshot_url == snapshot_url(timestamp, AREA_URL)
    assert response.archive_snapshot_timestamp == "2022-09-25T06:06:43Z"
    # The final URL cites the official council page, not the archive host.
    assert response.final_url == AREA_URL

    # The standard diagnostic accepts the archived copy as a valid official
    # result page and carries the capture citation into the audit record.
    result = fetch_and_diagnose_official_page(
        AREA_URL,
        WaybackOfficialPageClient(http),
    )
    assert (
        result.diagnostic.classification
        is OfficialPageClassification.VALID_ELECTION_RESULT_PAGE
    )
    assert result.diagnostic.retrieval_source == ARCHIVED_OFFICIAL_COPY
    assert result.diagnostic.archive_snapshot_url == snapshot_url(
        timestamp, AREA_URL
    )


def test_wayback_page_client_skips_captures_of_the_protection_stub() -> None:
    """A crawler can archive the challenge page; those captures carry no data.

    The fast path finds the newest capture but it stores the stub, so the
    client must fall back to the CDX index and serve an older usable capture.
    """

    newest = "20250101000000"
    usable = "20190720030604"
    http = FakeWaybackHTTP(
        cdx={AREA_URL: [[newest, AREA_URL], [usable, AREA_URL]]},
        pages={
            snapshot_url(newest, AREA_URL): (200, INCAPSULA_BODY),
            snapshot_url(usable, AREA_URL): (200, official_page_body("Addlestone")),
        },
        replay={
            snapshot_url("20991231235959", AREA_URL): (
                snapshot_url(newest, AREA_URL),
                200,
                INCAPSULA_BODY,
            )
        },
    )

    response = WaybackOfficialPageClient(http).fetch(AREA_URL)

    assert response.archive_snapshot_timestamp == "2019-07-20T03:06:04Z"


def test_wayback_page_client_raises_when_no_usable_capture_exists() -> None:
    """Probes that report no capture must not trigger slow CDX queries."""

    http = FakeWaybackHTTP(cdx={}, pages={})

    with pytest.raises(ArchivedSnapshotNotFoundError):
        WaybackOfficialPageClient(http).fetch(AREA_URL)

    # Every request was a fast-path probe; the CDX endpoint was never queried.
    assert all(
        not request.startswith(WAYBACK_CDX_ENDPOINT) for request in http.requests
    )


def test_fallback_client_returns_the_live_page_when_it_is_served() -> None:
    """A working live page needs no archive lookup at all."""

    live = CountingLiveClient(
        OfficialPageResponse(200, AREA_URL, official_page_body("Addlestone"))
    )
    http = FakeWaybackHTTP(cdx={}, pages={})
    client = ArchiveFallbackOfficialPageClient(live, WaybackOfficialPageClient(http))

    response = client.fetch(AREA_URL)

    assert response.retrieval_source == "live_official_page"
    assert live.calls == 1
    # No CDX or snapshot request was needed.
    assert http.requests == []


def test_fallback_client_switches_to_archive_and_stops_contacting_live_site() -> None:
    timestamp = "20220925060643"
    live = protected_live_page_client()
    http = FakeWaybackHTTP(
        cdx={},
        pages={},
        replay={
            snapshot_url("20991231235959", AREA_URL): (
                snapshot_url(timestamp, AREA_URL),
                200,
                official_page_body("Addlestone"),
            )
        },
    )
    client = ArchiveFallbackOfficialPageClient(live, WaybackOfficialPageClient(http))

    first = client.fetch(AREA_URL)
    second = client.fetch(AREA_URL)

    assert first.retrieval_source == ARCHIVED_OFFICIAL_COPY
    assert second.retrieval_source == ARCHIVED_OFFICIAL_COPY
    # After one protection response the client must not keep re-requesting the
    # blocked site for every remaining ward.
    assert live.calls == 1


def test_fallback_client_reports_protection_when_archive_has_no_capture() -> None:
    live = protected_live_page_client()
    http = FakeWaybackHTTP(cdx={}, pages={})
    client = ArchiveFallbackOfficialPageClient(live, WaybackOfficialPageClient(http))

    result = fetch_and_diagnose_official_page(AREA_URL, client)

    # The honest outcome is the live protection observation, so the workflow
    # falls back to indexed-only mode exactly as before this module existed.
    assert (
        result.diagnostic.classification
        is OfficialPageClassification.PROTECTION_PAGE
    )


def test_fallback_archive_client_serves_index_pages_from_captures() -> None:
    timestamp = "20210601000000"
    index_body = """
    <html><body>
      <h1>County Council Election 2021 - Thursday, 6 May 2021</h1>
      <a href="mgElectionAreaResults.aspx?ID=901&RPID=1">Addlestone</a>
    </body></html>
    """
    live = CountingLiveArchiveClient(
        OfficialArchiveResponse(200, UNCONFIGURED_INDEX_URL, INCAPSULA_BODY)
    )
    http = FakeWaybackHTTP(
        cdx={},
        pages={},
        replay={
            snapshot_url("20991231235959", UNCONFIGURED_INDEX_URL): (
                snapshot_url(timestamp, UNCONFIGURED_INDEX_URL),
                200,
                index_body,
            )
        },
    )
    client = ArchiveFallbackOfficialArchiveClient(
        live,
        WaybackOfficialArchiveClient(http),
    )

    response = client.fetch(UNCONFIGURED_INDEX_URL)

    assert response.status_code == 200
    assert response.final_url == UNCONFIGURED_INDEX_URL
    assert "Addlestone" in response.body
    assert live.calls == 1


def _end_to_end_wayback_fixture() -> FakeWaybackHTTP:
    """Model one small archived election: an index page plus two ward pages."""

    index_timestamp = "20210601000000"
    area_timestamp = "20220925060643"
    index_body = """
    <html><body>
      <h1>County Council Election 2021 - Thursday, 6 May 2021</h1>
      <a href="mgElectionAreaResults.aspx?ID=901&RPID=453691252">Addlestone</a>
      <a href="mgElectionAreaResults.aspx?ID=902&RPID=453691252">Guildford South-East</a>
    </body></html>
    """
    # The discovered result URLs keep the RPID display parameter published in
    # the archived index links; the fast path probes that exact URL first.
    area_901 = (
        "https://mycouncil.surreycc.gov.uk/"
        "mgElectionAreaResults.aspx?ID=901&RPID=453691252"
    )
    area_902 = (
        "https://mycouncil.surreycc.gov.uk/"
        "mgElectionAreaResults.aspx?ID=902&RPID=453691252"
    )
    hint = "20991231235959"
    return FakeWaybackHTTP(
        cdx={},
        pages={},
        replay={
            snapshot_url(hint, UNCONFIGURED_INDEX_URL): (
                snapshot_url(index_timestamp, UNCONFIGURED_INDEX_URL),
                200,
                index_body,
            ),
            snapshot_url(hint, area_901): (
                snapshot_url(area_timestamp, area_901),
                200,
                official_page_body("Addlestone"),
            ),
            snapshot_url(hint, area_902): (
                snapshot_url(area_timestamp, area_902),
                200,
                official_page_body("Guildford South-East"),
            ),
        },
    )


def test_workflow_completes_every_ward_from_archived_official_copies() -> None:
    """Blocked live site + archived official copies must yield Complete wards."""

    http = _end_to_end_wayback_fixture()
    page_client = ArchiveFallbackOfficialPageClient(
        protected_live_page_client(),
        WaybackOfficialPageClient(http),
    )
    archive_client = ArchiveFallbackOfficialArchiveClient(
        CountingLiveArchiveClient(
            OfficialArchiveResponse(200, UNCONFIGURED_INDEX_URL, INCAPSULA_BODY)
        ),
        WaybackOfficialArchiveClient(http),
    )
    provider = MockSearchProvider({})
    progress = []

    result = run_extraction_workflow(
        UNCONFIGURED_INDEX_URL,
        provider=provider,
        use_official_sources=True,
        require_indexed_search=True,
        official_page_client=page_client,
        official_archive_client=archive_client,
        progress_callback=progress.append,
    )

    # Scenario A acceptance shape: every discovered ward is Complete because
    # the archived official tables publish every required field.
    assert result.complete == 2
    assert result.incomplete == 0
    assert result.failed == 0
    assert {record.division_ward_name for record in result.records} == {
        "Addlestone",
        "Guildford South-East",
    }
    assert all(record.source_type.value == "official" for record in result.records)

    # The researcher is told that archived copies are in use.
    assert any(
        "archived official copies" in update.message for update in progress
    )

    # Every official attempt cites both the council URL and the exact capture.
    official_attempts = [
        attempt
        for attempt in result.extraction_attempts
        if attempt.source_type.value == "official"
    ]
    assert official_attempts
    for attempt in official_attempts:
        assert attempt.selected_urls[0].startswith(
            "https://mycouncil.surreycc.gov.uk/"
        )
        assert any(
            url.startswith("https://web.archive.org/web/")
            for url in attempt.selected_urls
        )

    # Field evidence carries the same citation on every extracted row.
    assert all(
        any(
            "archived official copy" in evidence.search_result_title
            for evidence in record.field_evidence
        )
        for record in result.records
    )

    # The workbook itself is complete and auditable.
    workbook = load_workbook(BytesIO(result.workbook_bytes), read_only=True)
    assert workbook.sheetnames[0] == "Index"
    assert "Addlestone" in workbook.sheetnames
    assert "Guildford South-East" in workbook.sheetnames
    assert "Extraction Log" in workbook.sheetnames
