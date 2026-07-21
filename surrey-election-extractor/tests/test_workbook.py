"""Unit tests for Surrey Election Extractor workbook generation."""

import re
from dataclasses import replace
from datetime import datetime, timezone

from openpyxl import load_workbook

from election_extractor.extraction import (
    CandidateResultRecord,
    EvidenceSourceType,
    ExtractionAttempt,
    ExtractionStatus,
)
from election_extractor.models import (
    DiscoveredElectionArea,
    DiscoveryStatus,
    ElectionStructureMetadata,
)
from election_extractor.validation import (
    PublishedVotingSummary,
    ValidationResult,
    ValidationStatus,
)
from election_extractor.workbook import generate_workbook


SOURCE_URL = (
    "https://mycouncil.surreycc.gov.uk/"
    "mgElectionAreaResults.aspx?ID=201&RPID=0"
)
SECOND_URL = (
    "https://mycouncil.surreycc.gov.uk/"
    "mgElectionAreaResults.aspx?ID=202&RPID=0"
)
SEAT_SOURCE_URL = "https://www.legislation.gov.uk/uksi/2012/1872/contents/made"


def candidate(
    name: str,
    party: str,
    votes: int | None,
    share: float | None,
    outcome: str,
    *,
    source_url: str = SOURCE_URL,
    ward: str = "Addlestone",
    missing_fields: tuple[str, ...] = (),
    status: ExtractionStatus = ExtractionStatus.COMPLETE,
) -> CandidateResultRecord:
    return CandidateResultRecord(
        election_name="2021 Surrey County Council election",
        election_date="6 May 2021",
        authority="Surrey County Council",
        division_ward_name=ward,
        number_of_seats=1,
        candidate_name=name,
        original_party_name=party,
        votes_received=votes,
        vote_share=share,
        outcome=outcome,
        electorate=10000,
        ballot_papers_issued=4000,
        ballot_papers_rejected=10,
        turnout=40.0,
        source_url=source_url,
        extraction_status=status,
        missing_fields=missing_fields,
    )


def validation(
    *,
    source_url: str = SOURCE_URL,
    ward: str = "Addlestone",
    status: ValidationStatus = ValidationStatus.PASSED,
    missing_fields: tuple[str, ...] = (),
    warnings: tuple[str, ...] = (),
    notes: tuple[str, ...] = ("Candidate totals checked.",),
    failed_checks: tuple[str, ...] = (),
) -> ValidationResult:
    return ValidationResult(
        election_name="2021 Surrey County Council election",
        election_year=2021,
        division_ward_name=ward,
        source_url=source_url,
        validation_status=status,
        fields_reviewed=("votes_received", "turnout"),
        failed_checks=failed_checks,
        warnings=warnings,
        validation_notes=notes,
        missing_fields=missing_fields,
        validation_timestamp=datetime(2026, 7, 14, tzinfo=timezone.utc),
        events=(),
    )


def discovery(
    *,
    source_url: str = SOURCE_URL,
    ward: str = "Addlestone",
) -> DiscoveredElectionArea:
    return DiscoveredElectionArea(
        election_year=2021,
        election_name="2021 Surrey County Council election",
        division_ward_name=ward,
        result_url=source_url,
        source_index_url=(
            "https://mycouncil.surreycc.gov.uk/"
            "mgElectionElectionAreaResults.aspx?EID=16"
        ),
        discovery_status=DiscoveryStatus.DISCOVERED,
    )


def attempt(
    *,
    source_url: str = SOURCE_URL,
    status: ExtractionStatus = ExtractionStatus.COMPLETE,
    error: str | None = None,
) -> ExtractionAttempt:
    return ExtractionAttempt(
        source_url=source_url,
        query=f'site:mycouncil.surreycc.gov.uk "{source_url}"',
        status=status,
        result_count=2,
        candidate_record_count=2,
        error=error,
    )


def structure_metadata(
    *,
    official_seats: int | None = None,
    secondary_seats: int | None = 1,
) -> ElectionStructureMetadata:
    """Create one separate, source-backed Seats metadata record for a test."""
    return ElectionStructureMetadata(
        election_year=2021,
        election_name="2021 Surrey County Council election",
        authority="Surrey County Council",
        division_or_ward_name="Ash",
        official_number_of_seats=official_seats,
        secondary_number_of_seats=secondary_seats,
        seat_source_type=(
            "The Surrey (Electoral Changes) Order 2012" if secondary_seats is not None else None
        ),
        seat_source_url=SEAT_SOURCE_URL if secondary_seats is not None else None,
        seat_evidence_text=(
            "Article 4 and the Schedule name Ash and provide for one councillor per division."
            if secondary_seats is not None
            else None
        ),
        confidence="High" if secondary_seats is not None else None,
        notes="Supplementary evidence only; official result-page Seats remains missing.",
    )


def standard_inputs():
    records = (
        candidate("Furey, John Raymond", "Conservative", 1146, 57.3, "Elected"),
        candidate("Example, Jane", "Labour and Co-operative", 854, 42.7, "Not elected"),
    )
    return (
        records,
        (validation(),),
        (discovery(),),
        (attempt(),),
        (PublishedVotingSummary(SOURCE_URL, total_votes=2000, valid_votes=2000),),
    )


def create_standard_workbook(tmp_path):
    path = tmp_path / "surrey_results.xlsx"
    generate_workbook(path, *standard_inputs())
    return path, load_workbook(path)


def find_row(sheet, value: str) -> int:
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == value:
            return row
    raise AssertionError(f"{value!r} was not found")


def test_01_index_worksheet_created_with_filters_and_freeze_panes(tmp_path) -> None:
    _, workbook = create_standard_workbook(tmp_path)
    sheet = workbook["Index"]

    assert workbook.sheetnames[0] == "Index"
    assert tuple(cell.value for cell in sheet[1]) == (
        "Ward or Division",
        "Worksheet",
        "Election",
        "Election Date",
        "Source URL",
        "Status",
        "Missing Fields",
        "Validation Notes",
        "Search Attempts",
    )
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:I2"
    assert "IndexTable" in sheet.tables
    assert sheet["F2"].value == "Complete"


def test_02_ward_worksheet_created_with_required_metadata(tmp_path) -> None:
    _, workbook = create_standard_workbook(tmp_path)
    sheet = workbook["Addlestone"]

    assert sheet["A1"].value == "Addlestone Election Result"
    assert sheet["B2"].value == "2021 Surrey County Council election"
    assert sheet["B3"].value == "6 May 2021"
    assert sheet["B4"].value == "Addlestone"
    assert sheet["B5"].value == "Complete"


def test_03_candidate_results_table_preserves_published_values(tmp_path) -> None:
    _, workbook = create_standard_workbook(tmp_path)
    sheet = workbook["Addlestone"]
    header_row = find_row(sheet, "Candidate")

    assert tuple(sheet.cell(header_row, column).value for column in range(1, 6)) == (
        "Candidate",
        "Party",
        "Votes",
        "Vote Share",
        "Outcome",
    )
    assert sheet.cell(header_row + 1, 1).value == "Furey, John Raymond"
    assert sheet.cell(header_row + 1, 2).value == "Conservative"
    assert sheet.cell(header_row + 2, 2).value == "Labour and Co-operative"
    assert sheet.cell(header_row + 1, 3).value == 1146
    assert sheet.cell(header_row + 1, 3).number_format == "#,##0"
    assert sheet.cell(header_row + 1, 4).value == 57.3
    assert sheet.cell(header_row + 1, 4).number_format == '0.0"%"'
    assert len([table for table in sheet.tables.values() if table.name.startswith("CandidateResults")]) == 1


def test_04_voting_summary_table_has_exact_rows(tmp_path) -> None:
    _, workbook = create_standard_workbook(tmp_path)
    sheet = workbook["Addlestone"]
    header_row = find_row(sheet, "Detail")
    values = {
        sheet.cell(row, 1).value: sheet.cell(row, 2).value
        for row in range(header_row + 1, header_row + 7)
    }

    assert values == {
        "Seats": 1,
        "Total votes": 2000,
        "Electorate": 10000,
        "Ballot papers issued": 4000,
        "Ballot papers rejected": 10,
        "Turnout": 40.0,
    }
    assert sheet.cell(header_row + 6, 2).number_format == '0.0"%"'
    assert len([table for table in sheet.tables.values() if table.name.startswith("VotingSummary")]) == 1


def test_05_missing_numeric_values_remain_blank(tmp_path) -> None:
    path = tmp_path / "missing.xlsx"
    records = (
        replace(
            standard_inputs()[0][0],
            votes_received=None,
            turnout=None,
            missing_fields=("votes_received", "turnout"),
            extraction_status=ExtractionStatus.INCOMPLETE,
        ),
    )
    generate_workbook(
        path,
        records,
        (
            validation(
                status=ValidationStatus.INCOMPLETE,
                missing_fields=("candidate[1].votes", "summary.total_votes", "summary.turnout"),
            ),
        ),
        (discovery(),),
        (attempt(status=ExtractionStatus.INCOMPLETE),),
        (PublishedVotingSummary(SOURCE_URL, total_votes=None, valid_votes=None),),
    )
    workbook = load_workbook(path)
    sheet = workbook["Addlestone"]
    candidate_header = find_row(sheet, "Candidate")
    summary_header = find_row(sheet, "Detail")

    assert sheet.cell(candidate_header + 1, 3).value is None
    summary = {
        sheet.cell(row, 1).value: sheet.cell(row, 2).value
        for row in range(summary_header + 1, summary_header + 7)
    }
    assert summary["Total votes"] is None
    assert summary["Turnout"] is None
    assert all(value != "Not found" for value in summary.values())


def test_06_invalid_worksheet_characters_are_removed(tmp_path) -> None:
    path = tmp_path / "invalid_name.xlsx"
    ward = "Long:Ward/Name?With*Invalid[Characters]Beyond Thirty One"
    record = candidate("Candidate", "Independent", 10, 100.0, "Elected", ward=ward)
    generate_workbook(
        path,
        (record,),
        (validation(ward=ward),),
        (discovery(ward=ward),),
        (attempt(),),
        (PublishedVotingSummary(SOURCE_URL, total_votes=10, valid_votes=10),),
    )
    workbook = load_workbook(path)
    ward_sheet_name = workbook.sheetnames[1]

    assert len(ward_sheet_name) <= 31
    assert not re.search(r"[:\\/?*\[\]]", ward_sheet_name)
    assert "LongWardName" in ward_sheet_name


def test_07_duplicate_worksheet_names_are_made_unique(tmp_path) -> None:
    path = tmp_path / "duplicates.xlsx"
    records = (
        candidate("Candidate One", "Independent", 10, 100.0, "Elected"),
        candidate(
            "Candidate Two",
            "Independent",
            20,
            100.0,
            "Elected",
            source_url=SECOND_URL,
        ),
    )
    generate_workbook(
        path,
        records,
        (
            validation(),
            validation(source_url=SECOND_URL),
        ),
        (
            discovery(),
            discovery(source_url=SECOND_URL),
        ),
        (
            attempt(),
            attempt(source_url=SECOND_URL),
        ),
        (
            PublishedVotingSummary(SOURCE_URL, 10, 10),
            PublishedVotingSummary(SECOND_URL, 20, 20),
        ),
    )
    workbook = load_workbook(path)

    assert "Addlestone" in workbook.sheetnames
    assert "Addlestone (2)" in workbook.sheetnames


def test_08_external_and_internal_hyperlinks_are_created(tmp_path) -> None:
    _, workbook = create_standard_workbook(tmp_path)
    index = workbook["Index"]
    ward = workbook["Addlestone"]

    assert index["B2"].hyperlink.target == "#'Addlestone'!A1"
    assert index["E2"].hyperlink.target == SOURCE_URL
    assert ward["B6"].hyperlink.target == SOURCE_URL


def test_09_extraction_log_created_without_credentials(tmp_path) -> None:
    _, workbook = create_standard_workbook(tmp_path)
    sheet = workbook["Extraction Log"]

    assert tuple(cell.value for cell in sheet[1]) == (
        "Ward or Division",
        "Attempt",
        "Query",
        "Results Returned",
        "Fields Found",
        "Fields Missing",
        "Warning or Error",
    )
    assert sheet["A2"].value == "Addlestone"
    assert sheet["B2"].value == 1
    assert sheet["D2"].value == 2
    assert "candidate_name" in sheet["E2"].value
    assert "api_key" not in " ".join(str(cell.value) for row in sheet for cell in row).casefold()


def test_10_workbook_generation_does_not_modify_inputs(tmp_path) -> None:
    inputs = standard_inputs()
    original = tuple(tuple(items) for items in inputs)

    generate_workbook(tmp_path / "unchanged.xlsx", *inputs)

    assert inputs == original


def test_11_failed_area_has_reason_and_no_fake_candidate_rows(tmp_path) -> None:
    path = tmp_path / "failed.xlsx"
    generate_workbook(
        path,
        (),
        (
            validation(
                status=ValidationStatus.FAILED,
                failed_checks=("vote_share",),
                notes=("Extraction could not be validated.",),
            ),
        ),
        (discovery(),),
        (attempt(status=ExtractionStatus.SEARCH_FAILED, error="SearchError"),),
        (),
    )
    workbook = load_workbook(path)
    sheet = workbook["Addlestone"]
    candidate_header = find_row(sheet, "Candidate")

    assert sheet["B5"].value == "Failed"
    assert "SearchError" in sheet["B10"].value
    assert sheet.cell(candidate_header + 1, 1).value is None


def test_11b_official_records_ignore_the_mandatory_indexed_audit_attempt(
    tmp_path,
) -> None:
    """A source-complete official row must not be downgraded by a required but
    unrelated indexed audit query that legitimately found no Google results.

    The Streamlit app runs one mandatory exact-URL indexed search for every
    result URL (the supervisor's search-workflow requirement) in addition to
    reading the official page. Google rarely indexes these obscure council
    pages, so that mandatory query commonly returns zero results even when the
    official table supplied every field. Regression coverage for a bug where
    ``workbook._area_status`` (unlike ``workflow._area_status``) counted that
    unrelated indexed attempt and therefore marked an otherwise Complete,
    fully official-sourced ward as Incomplete in the exported workbook.
    """
    official_records = tuple(
        replace(record, source_type=EvidenceSourceType.OFFICIAL)
        for record in (
            candidate("Furey, John Raymond", "Conservative", 1146, 57.3, "Elected"),
            candidate(
                "Example, Jane", "Labour and Co-operative", 854, 42.7, "Not elected"
            ),
        )
    )
    attempts = (
        # The mandatory exact-URL indexed audit query: it ran, found nothing,
        # and must not by itself downgrade the official-sourced result.
        ExtractionAttempt(
            source_url=SOURCE_URL,
            query=f'site:mycouncil.surreycc.gov.uk "{SOURCE_URL}"',
            status=ExtractionStatus.NO_EVIDENCE,
            result_count=0,
            candidate_record_count=0,
            source_type=EvidenceSourceType.INDEXED_SEARCH,
        ),
        # The official-page fetch that actually supplied the exported rows.
        ExtractionAttempt(
            source_url=SOURCE_URL,
            query=f"GET {SOURCE_URL}",
            status=ExtractionStatus.COMPLETE,
            result_count=1,
            candidate_record_count=2,
            source_type=EvidenceSourceType.OFFICIAL,
        ),
    )
    path = tmp_path / "official_only.xlsx"
    generate_workbook(
        path,
        official_records,
        (validation(),),
        (discovery(),),
        attempts,
        (PublishedVotingSummary(SOURCE_URL, total_votes=2000, valid_votes=2000),),
    )

    workbook = load_workbook(path)
    index_sheet = workbook["Index"]
    status_row = find_row(index_sheet, "Addlestone")
    assert index_sheet.cell(status_row, 6).value == "Complete"
    assert workbook["Addlestone"]["B5"].value == "Complete"


def test_12_election_structure_metadata_is_separate_and_source_backed(tmp_path) -> None:
    path = tmp_path / "structure_metadata.xlsx"
    records, validations, discoveries, attempts, summaries = standard_inputs()
    incomplete_records = (
        replace(
            records[0],
            number_of_seats=None,
            missing_fields=("number_of_seats",),
            extraction_status=ExtractionStatus.INCOMPLETE,
        ),
    )
    incomplete_validations = (
        validation(
            status=ValidationStatus.INCOMPLETE,
            missing_fields=("summary.number_of_seats",),
        ),
    )

    generate_workbook(
        path,
        incomplete_records,
        incomplete_validations,
        discoveries,
        attempts,
        summaries,
        election_structure_metadata=(structure_metadata(),),
    )
    workbook = load_workbook(path)
    sheet = workbook["Election Structure Metadata"]

    assert tuple(cell.value for cell in sheet[1]) == (
        "Election Year",
        "Election",
        "Authority",
        "Division or Ward",
        "Official Seats",
        "Secondary Seats",
        "Source Type",
        "Source URL",
        "Evidence",
        "Confidence",
        "Notes",
    )
    assert sheet["D2"].value == "Ash"
    assert sheet["E2"].value is None
    assert sheet["F2"].value == 1
    assert sheet["H2"].hyperlink.target == SEAT_SOURCE_URL
    assert "Article 4" in sheet["I2"].value
    assert "ElectionStructureMetadataTable" in sheet.tables

    # The presence of supplementary evidence must not change the official
    # extraction status or fill the official Seats cell in the ward worksheet.
    assert workbook["Index"]["F2"].value == "Incomplete"
    ward_sheet = workbook["Addlestone"]
    summary_header = find_row(ward_sheet, "Detail")
    assert ward_sheet.cell(summary_header + 1, 2).value is None
