"""Excel workbook generation for validated Surrey election results."""

import re
from collections.abc import Iterable, Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from election_extractor.extraction import (
    CandidateResultRecord,
    EvidenceSourceType,
    ExtractionAttempt,
    ExtractionStatus,
)
from election_extractor.models import DiscoveredElectionArea, ElectionStructureMetadata
from election_extractor.validation import (
    PublishedVotingSummary,
    ValidationResult,
    ValidationStatus,
)


# These headings are fixed by the workbook specification. Keeping them in one
# place prevents accidental wording changes between generation and tests.
INDEX_COLUMNS = (
    "Ward or Division",
    "Worksheet",
    "Election",
    "Election Date",
    "Source URL",
    "Status",
    "Missing Fields",
    "Validation Notes",
    "Search Attempts",
)

CANDIDATE_COLUMNS = ("Candidate", "Party", "Votes", "Vote Share", "Outcome")
SUMMARY_ROWS = (
    ("Seats", "number_of_seats"),
    ("Total votes", "total_votes"),
    ("Electorate", "electorate"),
    ("Ballot papers issued", "ballot_papers_issued"),
    ("Ballot papers rejected", "ballot_papers_rejected"),
    ("Turnout", "turnout"),
)
LOG_COLUMNS = (
    "Ward or Division",
    "Attempt",
    "Query",
    "Results Returned",
    "Fields Found",
    "Fields Missing",
    "Warning or Error",
)
ELECTION_STRUCTURE_METADATA_COLUMNS = (
    "Election Year",
    "Election",
    "Authority",
    "Division or Ward",
    "Official Seats",
    "Secondary Seats",
    "Source Type",
    "Source URL",
    "Evidence",
    "Confidence",
    "Notes",
)

INVALID_SHEET_CHARACTERS = re.compile(r"[:\\/?*\[\]]")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
INCOMPLETE_FILL = PatternFill("solid", fgColor="FFF2CC")
FAILED_FILL = PatternFill("solid", fgColor="F4CCCC")
WHITE_BOLD_FONT = Font(color="FFFFFF", bold=True)
LINK_FONT = Font(color="0563C1", underline="single")
LABEL_FONT = Font(bold=True)


def _is_missing(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _first_non_missing(values: Iterable[object]) -> object | None:
    for value in values:
        if not _is_missing(value):
            return value
    return None


def _consensus(values: Iterable[object]) -> object | None:
    """Return a shared value or blank when published values conflict."""
    distinct: dict[tuple[type[object], str], object] = {}
    for value in values:
        if _is_missing(value):
            continue
        key = (type(value), str(value).casefold())
        distinct.setdefault(key, value)
    return next(iter(distinct.values())) if len(distinct) == 1 else None


def _join_text(values: Iterable[str | None]) -> str:
    """Join unique non-empty audit text without changing its wording."""
    output = []
    seen = set()
    for value in values:
        if not value:
            continue
        cleaned = value.strip()
        if cleaned and cleaned not in seen:
            seen.add(cleaned)
            output.append(cleaned)
    return " | ".join(output)


def safe_worksheet_name(name: str | None, used_names: set[str]) -> str:
    """Create a recognisable, valid and case-insensitively unique sheet name."""
    # Excel rejects the characters matched here and limits names to 31
    # characters. The numeric suffix keeps repeated ward names recognisable.
    cleaned = INVALID_SHEET_CHARACTERS.sub("", name or "")
    cleaned = " ".join(cleaned.split()).strip("'") or "Ward"
    base = cleaned[:31]
    candidate = base
    suffix_number = 2
    used_casefold = {item.casefold() for item in used_names}
    while candidate.casefold() in used_casefold:
        suffix = f" ({suffix_number})"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        suffix_number += 1
    used_names.add(candidate)
    return candidate


def _external_hyperlink(cell: Cell, url: str | None) -> None:
    if not url:
        return
    cell.value = url
    cell.hyperlink = url
    cell.font = LINK_FONT


def _internal_hyperlink(cell: Cell, sheet_name: str) -> None:
    escaped_name = sheet_name.replace("'", "''")
    cell.value = sheet_name
    cell.hyperlink = f"#'{escaped_name}'!A1"
    cell.font = LINK_FONT


def _table_style(table: Table) -> None:
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )


def _style_header(cells: Iterable[Cell]) -> None:
    for cell in cells:
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_section_title(sheet: Worksheet, row: int, title: str, width: int) -> None:
    sheet.cell(row=row, column=1, value=title)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    cell = sheet.cell(row=row, column=1)
    cell.fill = SECTION_FILL
    cell.font = Font(bold=True, size=12)


def _autosize_columns(sheet: Worksheet, minimum: int = 10, maximum: int = 55) -> None:
    for column_cells in sheet.columns:
        letter = column_cells[0].column_letter
        longest = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=0,
        )
        sheet.column_dimensions[letter].width = min(max(longest + 2, minimum), maximum)


def _configure_print_layout(sheet: Worksheet) -> None:
    """Keep wide audit tables together in spreadsheet print previews."""
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4


def _percentage_format(value: float | int | None) -> str:
    """Display only precision supported by the extracted numeric value."""
    if value is None:
        return "General"
    text = str(value)
    decimals = len(text.partition(".")[2])
    return f'0.{"0" * decimals}"%"' if decimals else '0"%"'


def _group_by_source_url(items: Sequence[object]) -> dict[str, list[object]]:
    """Group records from different stages around the same official result URL."""
    groups: dict[str, list[object]] = {}
    for item in items:
        source_url = getattr(item, "source_url", None) or getattr(item, "result_url", None)
        if source_url:
            groups.setdefault(source_url, []).append(item)
    return groups


def _area_status(
    records: Sequence[CandidateResultRecord],
    validations: Sequence[ValidationResult],
    attempts: Sequence[ExtractionAttempt],
) -> str:
    """Map extraction and validation outcomes to the three required text statuses."""
    # A confirmed search or validation failure takes precedence over all other
    # evidence so the workbook never presents a failed area as usable data.
    if any(attempt.status is ExtractionStatus.SEARCH_FAILED for attempt in attempts):
        return "Failed"
    if any(result.validation_status is ValidationStatus.FAILED for result in validations):
        return "Failed"
    if any(record.extraction_status is ExtractionStatus.SEARCH_FAILED for record in records):
        return "Failed"
    # A result URL with no reliable candidate rows meets the supervisor's
    # explicit Failed definition, even if the search itself returned normally.
    if not records:
        return "Failed"

    # Bug fix: this function used to count the mandatory exact-URL indexed
    # audit query as an ordinary attempt. That query almost never gets a
    # Google hit for these obscure council pages, so every ward that was
    # actually Complete from the official table was still exported as
    # Incomplete, even though the app's own progress bar said Complete. The
    # fix is the same one already applied in workflow._area_status: if every
    # exported record came from the official page, only official-tier
    # attempts should count towards completeness.
    record_source_types = {record.source_type for record in records}
    status_attempts = attempts
    if record_source_types == {EvidenceSourceType.OFFICIAL}:
        status_attempts = tuple(
            attempt
            for attempt in attempts
            if attempt.source_type is EvidenceSourceType.OFFICIAL
        )

    # Warnings, missing validation, missing candidates or partial attempts all
    # remain visible as Incomplete instead of being silently treated as success.
    extraction_incomplete = any(
        record.extraction_status is not ExtractionStatus.COMPLETE for record in records
    ) or any(
        attempt.status is not ExtractionStatus.COMPLETE for attempt in status_attempts
    )
    validation_incomplete = not validations or any(
        result.validation_status in {ValidationStatus.INCOMPLETE, ValidationStatus.WARNING}
        for result in validations
    )
    if extraction_incomplete or validation_incomplete:
        return "Incomplete"
    return "Complete"


def _fields_found(records: Sequence[CandidateResultRecord]) -> str:
    fields = (
        "election_name",
        "election_date",
        "authority",
        "division_ward_name",
        "number_of_seats",
        "candidate_name",
        "original_party_name",
        "votes_received",
        "vote_share",
        "outcome",
        "electorate",
        "ballot_papers_issued",
        "ballot_papers_rejected",
        "turnout",
    )
    return ", ".join(
        field
        for field in fields
        if any(not _is_missing(getattr(record, field)) for record in records)
    )


def _missing_fields(
    records: Sequence[CandidateResultRecord],
    validations: Sequence[ValidationResult],
) -> str:
    fields = []
    for record in records:
        fields.extend(record.missing_fields)
    for validation in validations:
        fields.extend(validation.missing_fields)
    return _join_text(fields)


def _validation_notes(validations: Sequence[ValidationResult]) -> str:
    notes = []
    for validation in validations:
        notes.extend(validation.validation_notes)
        notes.extend(validation.warnings)
    return _join_text(notes)


def _failure_reason(
    validations: Sequence[ValidationResult],
    attempts: Sequence[ExtractionAttempt],
) -> str:
    reasons = [attempt.error for attempt in attempts if attempt.error]
    for validation in validations:
        if validation.failed_checks:
            reasons.append(f"Failed validation checks: {', '.join(validation.failed_checks)}")
    return _join_text(reasons)


def _attempt_text(attempts: Sequence[ExtractionAttempt]) -> str:
    lines = []
    for number, attempt in enumerate(attempts, start=1):
        text = (
            f"Attempt {number}: {attempt.status.value}; results={attempt.result_count}; "
            f"query={attempt.query}"
        )
        if attempt.error:
            text += f"; error={attempt.error}"
        lines.append(text)
    return "\n".join(lines)


def _summary_value(
    field_name: str,
    records: Sequence[CandidateResultRecord],
    summaries: Sequence[PublishedVotingSummary],
) -> object | None:
    """Read published summary values without calculating replacements."""
    # Total votes belongs to the published summary model. Other values are
    # repeated on candidate records and are written only when they agree.
    if field_name == "total_votes":
        return _consensus(summary.total_votes for summary in summaries)
    return _consensus(getattr(record, field_name) for record in records)


def _write_ward_sheet(
    sheet: Worksheet,
    table_number: int,
    ward_name: str,
    election_name: str | None,
    election_date: str | None,
    source_url: str,
    status: str,
    records: Sequence[CandidateResultRecord],
    validations: Sequence[ValidationResult],
    attempts: Sequence[ExtractionAttempt],
    summaries: Sequence[PublishedVotingSummary],
) -> None:
    """Write metadata, candidate results and voting summary for one area."""
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    _configure_print_layout(sheet)
    sheet.merge_cells("A1:E1")
    sheet["A1"] = f"{ward_name} Election Result"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A1"].fill = SECTION_FILL

    # Audit information stays above the data tables so incomplete and failed
    # extractions can be understood without opening the Extraction Log.
    metadata = (
        ("Election name", election_name or ""),
        ("Election date", election_date or ""),
        ("Ward/division name", ward_name),
        ("Extraction status", status),
        ("Source URL", source_url),
        ("Missing fields", _missing_fields(records, validations)),
        ("Validation notes", _validation_notes(validations)),
        ("Search attempts", _attempt_text(attempts)),
        ("Failure reason", _failure_reason(validations, attempts) if status == "Failed" else ""),
    )
    for row, (label, value) in enumerate(metadata, start=2):
        sheet.cell(row=row, column=1, value=label).font = LABEL_FONT
        value_cell = sheet.cell(row=row, column=2, value=value)
        value_cell.alignment = Alignment(vertical="top", wrap_text=True)
        if label == "Source URL":
            _external_hyperlink(value_cell, source_url)

    # The supervisor requires an explicit notice rather than relying only on a
    # status cell in the metadata block. Keep this statement outside numeric
    # cells so missing source values can remain genuinely blank.
    if status == "Incomplete":
        sheet.merge_cells("A11:E11")
        sheet["A11"] = (
            "Status: Incomplete — missing information was not inferred or invented."
        )
        sheet["A11"].font = Font(bold=True, color="9C5700")
        sheet["A11"].fill = INCOMPLETE_FILL
    elif status == "Failed":
        sheet.merge_cells("A11:E11")
        sheet["A11"] = (
            "Status: Failed — no reliable result data was fabricated or substituted."
        )
        sheet["A11"].font = Font(bold=True, color="9C0006")
        sheet["A11"].fill = FAILED_FILL

    # Row 12 remains blank between the audit notice and the first table.
    candidate_title_row = 13
    candidate_header_row = candidate_title_row + 1
    _style_section_title(
        sheet,
        candidate_title_row,
        f"{ward_name} Election Results",
        len(CANDIDATE_COLUMNS),
    )
    for column, heading in enumerate(CANDIDATE_COLUMNS, start=1):
        sheet.cell(row=candidate_header_row, column=column, value=heading)
    _style_header(sheet[candidate_header_row])

    # Candidate and party values are copied directly from extraction records.
    # No standardisation, recalculation or replacement is performed here.
    for row_offset, record in enumerate(records, start=1):
        row = candidate_header_row + row_offset
        values = (
            record.candidate_name,
            record.original_party_name,
            record.votes_received,
            record.vote_share,
            record.outcome,
        )
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value)
        sheet.cell(row=row, column=3).number_format = "#,##0"
        sheet.cell(row=row, column=4).number_format = _percentage_format(record.vote_share)

    candidate_end_row = candidate_header_row + len(records)
    # Do not create an Excel candidate table when extraction produced no real
    # candidates. This prevents failed areas from gaining fabricated rows.
    if records:
        candidate_table = Table(
            displayName=f"CandidateResults{table_number}",
            ref=f"A{candidate_header_row}:E{candidate_end_row}",
        )
        _table_style(candidate_table)
        sheet.add_table(candidate_table)

    # Two fully blank rows separate the candidate table from Voting Summary.
    summary_title_row = candidate_end_row + 3
    summary_header_row = summary_title_row + 1
    _style_section_title(sheet, summary_title_row, "Voting Summary", 2)
    sheet.cell(row=summary_header_row, column=1, value="Detail")
    sheet.cell(row=summary_header_row, column=2, value="Number")
    _style_header(sheet[summary_header_row][:2])

    # Missing or conflicting numeric values are left as blank cells. Validation
    # output is explanatory only and never replaces a published source value.
    for offset, (label, field_name) in enumerate(SUMMARY_ROWS, start=1):
        row = summary_header_row + offset
        value = _summary_value(field_name, records, summaries)
        sheet.cell(row=row, column=1, value=label)
        value_cell = sheet.cell(row=row, column=2, value=value)
        if field_name == "turnout":
            value_cell.number_format = _percentage_format(value)
        else:
            value_cell.number_format = "#,##0"

    summary_end_row = summary_header_row + len(SUMMARY_ROWS)
    summary_table = Table(
        displayName=f"VotingSummary{table_number}",
        ref=f"A{summary_header_row}:B{summary_end_row}",
    )
    _table_style(summary_table)
    sheet.add_table(summary_table)

    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 16
    sheet.column_dimensions["E"].width = 20
    # Long URLs and audit messages need explicit row heights because openpyxl
    # cannot ask Excel to auto-fit wrapped rows while creating the workbook.
    sheet.row_dimensions[6].height = 42
    sheet.row_dimensions[7].height = 30
    sheet.row_dimensions[8].height = 45
    sheet.row_dimensions[9].height = 75
    sheet.row_dimensions[10].height = 30
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical="top",
                wrap_text=True,
            )
    for row in range(candidate_header_row + 1, candidate_end_row + 1):
        sheet.cell(row=row, column=4).alignment = Alignment(
            horizontal="right", vertical="top", wrap_text=True
        )
        sheet.cell(row=row, column=5).alignment = Alignment(
            horizontal="left", vertical="top", wrap_text=True, indent=1
        )
    sheet.print_area = f"A1:E{summary_end_row}"


def _write_election_structure_metadata_sheet(
    sheet: Worksheet,
    metadata_records: Sequence[ElectionStructureMetadata],
) -> None:
    """Write supplementary structure evidence without changing official values."""
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(ELECTION_STRUCTURE_METADATA_COLUMNS)
    _style_header(sheet[1])

    # Official and supplementary Seats values occupy different cells.  This
    # makes the source boundary visible in every exported workbook and avoids
    # presenting a secondary value as if Surrey published it on the result page.
    for row_number, metadata in enumerate(metadata_records, start=2):
        sheet.append(
            (
                metadata.election_year,
                metadata.election_name,
                metadata.authority,
                metadata.division_or_ward_name,
                metadata.official_number_of_seats,
                metadata.secondary_number_of_seats,
                metadata.seat_source_type,
                metadata.seat_source_url,
                metadata.seat_evidence_text,
                metadata.confidence,
                metadata.notes,
            )
        )
        # The source URL is a link for auditability, while its text remains
        # visible to users who view the worksheet without following hyperlinks.
        _external_hyperlink(sheet.cell(row=row_number, column=8), metadata.seat_source_url)

    if metadata_records:
        table = Table(
            displayName="ElectionStructureMetadataTable",
            ref=f"A1:K{len(metadata_records) + 1}",
        )
        _table_style(table)
        sheet.add_table(table)
        sheet.auto_filter.ref = f"A1:K{len(metadata_records) + 1}"

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in (5, 6):
        for row in range(2, len(metadata_records) + 2):
            sheet.cell(row=row, column=column).number_format = "#,##0"

    _autosize_columns(sheet, minimum=12, maximum=55)
    sheet.column_dimensions["H"].width = 55
    sheet.column_dimensions["I"].width = 55
    sheet.column_dimensions["K"].width = 55
    _configure_print_layout(sheet)
    sheet.print_area = f"A1:K{max(len(metadata_records) + 1, 1)}"


def generate_workbook(
    output_path: str | Path,
    records: Sequence[CandidateResultRecord],
    validation_results: Sequence[ValidationResult],
    discovery_areas: Sequence[DiscoveredElectionArea] = (),
    extraction_attempts: Sequence[ExtractionAttempt] = (),
    voting_summaries: Sequence[PublishedVotingSummary] = (),
    election_structure_metadata: Sequence[ElectionStructureMetadata] = (),
) -> Path:
    """Generate one workbook without changing extracted or validated inputs.

    Each official result URL acts as the stable key joining discovery,
    extraction, validation and audit information for one ward or division.
    """
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    # Build read-only lookup groups; the supplied dataclass records are never
    # mutated while the workbook is assembled.
    records_by_url = _group_by_source_url(records)
    validations_by_url = _group_by_source_url(validation_results)
    discoveries_by_url = _group_by_source_url(discovery_areas)
    attempts_by_url = _group_by_source_url(extraction_attempts)
    summaries_by_url = _group_by_source_url(voting_summaries)
    # Use the union so failed or incomplete areas still receive an Index row and
    # worksheet even when one processing stage produced no records.
    source_urls = sorted(
        set(records_by_url)
        | set(validations_by_url)
        | set(discoveries_by_url)
        | set(attempts_by_url)
        | set(summaries_by_url)
    )

    # openpyxl creates one default sheet. Reusing it guarantees Index remains
    # the first worksheet as required.
    workbook = Workbook()
    index_sheet = workbook.active
    index_sheet.title = "Index"
    index_sheet.sheet_view.showGridLines = False
    index_sheet.freeze_panes = "A2"
    index_sheet.append(INDEX_COLUMNS)
    _style_header(index_sheet[1])

    used_names = {"Index", "Extraction Log"}
    if election_structure_metadata:
        # Reserve this name only when the optional research metadata worksheet
        # will actually be written. The Streamlit prompt otherwise requires the
        # workbook to contain only Index, area sheets and Extraction Log.
        used_names.add("Election Structure Metadata")
    area_rows = []
    ward_sheets = []
    for table_number, source_url in enumerate(source_urls, start=1):
        area_records = tuple(records_by_url.get(source_url, ()))
        area_validations = tuple(validations_by_url.get(source_url, ()))
        area_discoveries = tuple(discoveries_by_url.get(source_url, ()))
        area_attempts = tuple(attempts_by_url.get(source_url, ()))
        area_summaries = tuple(summaries_by_url.get(source_url, ()))

        # Prefer discovery metadata, then validation, then extraction. This lets
        # a failed extraction retain the ward name discovered in the earlier stage.
        ward_name = str(
            _first_non_missing(
                [
                    *(area.division_ward_name for area in area_discoveries),
                    *(result.division_ward_name for result in area_validations),
                    *(record.division_ward_name for record in area_records),
                ]
            )
            or "Unnamed Ward"
        )
        election_name = _first_non_missing(
            [
                *(area.election_name for area in area_discoveries),
                *(result.election_name for result in area_validations),
                *(record.election_name for record in area_records),
            ]
        )
        election_date = _first_non_missing(record.election_date for record in area_records)
        status = _area_status(area_records, area_validations, area_attempts)
        worksheet_name = safe_worksheet_name(ward_name, used_names)
        ward_sheet = workbook.create_sheet(worksheet_name)
        _write_ward_sheet(
            ward_sheet,
            table_number,
            ward_name,
            str(election_name) if election_name is not None else None,
            str(election_date) if election_date is not None else None,
            source_url,
            status,
            area_records,
            area_validations,
            area_attempts,
            area_summaries,
        )
        ward_sheets.append((source_url, ward_name, worksheet_name, area_records, area_validations, area_attempts))

        area_rows.append(
            (
                ward_name,
                worksheet_name,
                election_name,
                election_date,
                source_url,
                status,
                _missing_fields(area_records, area_validations),
                _validation_notes(area_validations),
                len(area_attempts),
            )
        )

    # Index provides both an internal link to the area sheet and an external
    # link back to the published source page.
    for row_number, values in enumerate(area_rows, start=2):
        index_sheet.append(values)
        _internal_hyperlink(index_sheet.cell(row=row_number, column=2), str(values[1]))
        _external_hyperlink(index_sheet.cell(row=row_number, column=5), str(values[4]))

    if area_rows:
        index_table = Table(displayName="IndexTable", ref=f"A1:I{len(area_rows) + 1}")
        _table_style(index_table)
        index_sheet.add_table(index_table)
        index_sheet.auto_filter.ref = f"A1:I{len(area_rows) + 1}"
    _autosize_columns(index_sheet)
    _configure_print_layout(index_sheet)
    for row in index_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal,
                vertical="top",
                wrap_text=True,
            )
    index_sheet.print_area = f"A1:I{max(len(area_rows) + 1, 1)}"

    # The log is built only from non-sensitive attempt fields. Provider keys and
    # authentication data are not accepted by this workbook interface.
    log_sheet = workbook.create_sheet("Extraction Log")
    log_sheet.sheet_view.showGridLines = False
    log_sheet.freeze_panes = "A2"
    log_sheet.append(LOG_COLUMNS)
    _style_header(log_sheet[1])
    log_row = 2
    for _, ward_name, _, area_records, area_validations, area_attempts in ward_sheets:
        found = _fields_found(area_records)
        missing = _missing_fields(area_records, area_validations)
        warning = _join_text(
            [
                _failure_reason(area_validations, area_attempts),
                *(
                    warning_text
                    for validation in area_validations
                    for warning_text in validation.warnings
                ),
            ]
        )
        for attempt_number, attempt in enumerate(area_attempts, start=1):
            log_sheet.append(
                (
                    ward_name,
                    attempt_number,
                    attempt.query,
                    attempt.result_count,
                    found,
                    missing,
                    attempt.error or warning,
                )
            )
            log_row += 1

    if log_row > 2:
        log_table = Table(displayName="ExtractionLogTable", ref=f"A1:G{log_row - 1}")
        _table_style(log_table)
        log_sheet.add_table(log_table)
        log_sheet.auto_filter.ref = f"A1:G{log_row - 1}"
    _autosize_columns(log_sheet)
    _configure_print_layout(log_sheet)
    for row in log_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    log_sheet.print_area = f"A1:G{max(log_row - 1, 1)}"

    if election_structure_metadata:
        # This optional research worksheet remains independent of area status
        # and validation. Supplementary evidence must not turn an incomplete
        # official extraction into a complete one. It is omitted from ordinary
        # Streamlit exports because that workflow supplies no such records.
        metadata_sheet = workbook.create_sheet("Election Structure Metadata")
        _write_election_structure_metadata_sheet(metadata_sheet, election_structure_metadata)

    workbook.save(output)
    return output
