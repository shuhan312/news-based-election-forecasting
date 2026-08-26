"""Build the reform_uk sub-field validation workbook: the 14 D4 articles
the model marked applicable=true, plus 8 controls it marked false, hash-
shuffled together so the reviewer cannot tell which is which. Blind by
construction: no model output appears on the sheet.

Why controls at all: scoring only the model's positives would hand the
human an all-yes applicable column, a constant marginal, and a kappa
that collapses to 0 regardless of real agreement - the prevalence
problem already documented in eligibility_manual_review_methodology.md
section 8.1. Controls are drawn in sha256(article_id) order from the
applicable=false pool, reform-mentioning ones first (they are the near
misses that actually test the boundary), so the draw is reproducible
and unsteerable."""

import csv
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

csv.field_size_limit(10_000_000)
REPO = Path("/Users/sl1425/irp-sl1425")
OUT_DIR = REPO / "outputs" / "d4_validation_labelling_2026-07-29"
OUT = OUT_DIR / "D4_ReformUK_Subfields_Labelling.xlsx"

N_CONTROLS = 8


def hash_order(ids):
    return sorted(ids, key=lambda a: hashlib.sha256(a.encode()).hexdigest())


ARIAL = "Arial"
F_HDR = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_BODY = Font(name=ARIAL, size=10)
F_TITLE = Font(name=ARIAL, size=13, bold=True)
FILL_HDR = PatternFill("solid", fgColor="1F4E79")
FILL_YELLOW = PatternFill("solid", fgColor="FFFF00")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def main() -> None:
    llm = json.loads((REPO / "llm_context/d4_llm_outputs.json").read_text())
    sample = {r["article_id"]: r for r in csv.DictReader(
        open(REPO / "llm_context/d4_validation_sample_v1.csv"))}

    applicable_true, false_mentioning, false_other = [], [], []
    for result in llm["layers"]["consequence"]:
        record = result.get("record")
        if not record:
            continue
        article_id = result["article_id"]
        if (record.get("reform_uk") or {}).get("applicable"):
            applicable_true.append(article_id)
        elif sample.get(article_id, {}).get("mentions_reform") == "true":
            false_mentioning.append(article_id)
        else:
            false_other.append(article_id)

    # Near misses test the applicable boundary. Top up from the remaining
    # false cases only when the mention pool is too small.
    controls = (hash_order(false_mentioning)
                + hash_order(false_other))[:N_CONTROLS]
    rows = sorted(applicable_true + controls,
                  key=lambda a: hashlib.sha256(a.encode()).hexdigest())
    print(f"{len(applicable_true)} applicable + {len(controls)} controls "
          f"= {len(rows)} rows")

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    for index, width in enumerate([3, 30, 108], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.cell(
        row=1,
        column=2,
        value="Reform UK sub-field validation — 22 articles × 6 fields",
    ).font = F_TITLE
    instructions = [
        ("Purpose", "The consequence layer failed its validation gate, but the separate reform_uk sub-field block built on it has not been validated. Some of these 22 articles concern Reform UK; others merely use the word 'reform'. Review each article independently for comparison with the model. If this block passes its gate, the Reform-specific features will be retained; otherwise they will be discarded."),
        ("ru_applicable", "Is the article substantively about Reform UK as a political party, rather than mentioning it in passing or using 'reform' in its ordinary sense? Select yes or no. If no, leave the remaining five fields blank."),
        ("ru_growth_suggested", "Does the article suggest that support for Reform UK is growing? Select yes or no."),
        ("ru_credible_challenger", "Does the article present Reform UK as a credible challenger rather than a fringe party? Select yes or no."),
        ("ru_established_support_affected", "Which established parties does the article suggest may lose support to Reform UK? Choose from conservative / labour / liberal_democrat, separated by commas; enter none if none apply."),
        ("ru_switching_directions", "Does the article mention voters switching to Reform UK? Choose from con_to_reform / lab_to_reform / ld_to_reform, separated by commas; enter none if none apply."),
        ("ru_signal_nature", "What is the nature of the signal? Choose from national_momentum / local_campaign_strength / protest_voting / anti_incumbent_sentiment, separated by commas; enter none if none apply."),
        ("Important", "Record only what the article states or implies. Volume of coverage must never be treated automatically as votes. If a boolean is uncertain, select no as the conservative judgement. Estimated completion time: 20 minutes."),
    ]
    row_number = 3
    for label, instruction in instructions:
        ws.cell(row=row_number, column=2, value=label).font = Font(
            name=ARIAL, size=11, bold=True)
        cell = ws.cell(row=row_number, column=3, value=instruction)
        cell.font = F_BODY
        cell.alignment = WRAP
        row_number += 2

    ws = wb.create_sheet("RU labels - 22 articles")
    headings = [
        "article_id", "election_id", "headline",
        "article excerpt (first 1,000 characters)", "full-text path",
        "ru_applicable", "ru_growth_suggested", "ru_credible_challenger",
        "ru_established_support_affected", "ru_switching_directions",
        "ru_signal_nature", "notes",
    ]
    for column, heading in enumerate(headings, 1):
        cell = ws.cell(row=1, column=column, value=heading)
        cell.font = F_HDR
        cell.fill = FILL_HDR
        cell.alignment = WRAP
    widths = [30, 13, 40, 70, 32, 13, 16, 18, 28, 26, 30, 24]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    headlines = {r["article_id"]: r["headline"] for r in csv.DictReader(
        open(REPO / "news_collection/full_corpus_review.csv"))}
    for row_number, article_id in enumerate(rows, start=2):
        text_path = REPO / "data/raw/news/text" / f"{article_id}.txt"
        excerpt = text_path.read_text(
            encoding="utf-8", errors="replace")[:1000]
        values = [
            article_id, sample[article_id]["election_id"],
            headlines.get(article_id, ""), excerpt, str(text_path),
        ] + [""] * 7
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row_number, column=column, value=value)
            cell.font = F_BODY
            cell.alignment = WRAP if column in (3, 4) else TOP
        for column in range(6, 12):
            ws.cell(row=row_number, column=column).fill = FILL_YELLOW

    validation = DataValidation(
        type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(validation)
    for column in ("F", "G", "H"):
        validation.add(f"{column}2:{column}{len(rows) + 1}")
    ws.freeze_panes = "C2"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.active = 0
    wb.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    main()
