"""Convert the validated 2013 SerpAPI-extractor workbook into
data/elections/2013_scc_results.csv, in the same column shape as the
completed historical baseline table (data/elections/results_2017_2024.csv),
closing the biggest remaining gap in the Raw News Corpus Construction
stage: 2013 currently has ZERO ward-tier local-news queries because no
committed division/candidate table for 2013 has ever existed in
data/elections/ (only official_scc_candidate_results_test.csv, a
13-row test fixture covering a single division).

Source of truth
----------------
surrey-election-extractor/outputs/2013_full_extraction/
    surrey_county_council_2013.xlsx
already exists, produced by the same validated SerpAPI-based pipeline
used for 2026 (src/convert_2026_extractor_output.py): 81/81 divisions
discovered and extracted, 358 candidate records, 0 failed extraction
attempts, 0 failed validation checks. This script only READS that
workbook and reshapes it; it does not re-run any extraction.

What is genuinely missing from the 2013 official pages (verified by
inspecting every ward's Index row, not by trusting a summary count)
------------------------------------------------------------------
Every one of the 81 wards is marked "Incomplete" in the source
workbook, but for a uniform, well-understood reason across 80 of them:
the 2013 official result pages never publish "ballot papers issued" or
a turnout percentage at all (confirmed directly: the Voting Summary
table for a sample ward has "Ballot papers issued" and "Turnout" rows
present with blank values, while Seats, Total votes, Electorate and
Ballot papers rejected are all populated). This is not a parsing
failure - the source page genuinely does not carry the figure - and
matches something already decided for this project: turnout for 2013
may be sourced from Wikipedia once cross-validated against at least 10
official-page examples (see project history/meeting notes). This
script does NOT perform that cross-validation or pull in Wikipedia
data - it leaves turnout_percent and people_who_voted blank for 2013,
exactly as the source leaves them, and flags every row so a follow-up
script can join in the already-approved Wikipedia figures later
without this script silently inventing anything in the meantime.

The remaining ward (1 of 81) is additionally missing one candidate's
party affiliation entirely (not published on the source page) - left
blank in party_raw/party_canonical rather than guessed at, and logged.

Party standardisation
----------------------
Party names are matched against the already-completed
data/elections/party_name_standardisation.csv lookup where possible.
Unmatched strings are kept as the raw published string (never
invented) and flagged in the conversion log, exactly as
convert_2026_extractor_output.py does, so any genuinely new 2013-era
party label can be added to the standardisation table deliberately.

Usage:
    python3 src/convert_2013_extractor_output.py
"""

import csv
from pathlib import Path

import openpyxl

WORKBOOK = Path("surrey-election-extractor/outputs/2013_full_extraction/"
                "surrey_county_council_2013.xlsx")
COUNCIL = "Surrey County Council"
POLLING_DATE = "2013-05-02"
OUT = Path("data/elections/2013_scc_results.csv")
LOG_OUT = Path("data/elections/2013_conversion_log.csv")
PARTY_LOOKUP_PATH = Path("data/elections/party_name_standardisation.csv")

# Same column shape as results_2017_2024.csv, so 2013 slots into the
# existing feature pipeline (scc_ward_data_2017_2021-style readers)
# without a bespoke schema. turnout_percent and people_who_voted are
# left blank for every row - genuinely absent from the source, not
# computed or guessed - pending the already-approved Wikipedia
# cross-validation workflow (see module docstring). total_candidate_votes
# and ward_index_status are appended, matching the 2026 converter's
# convention, for the same audit reasons.
OUTPUT_COLUMNS = [
    "year", "council", "ward", "party_raw", "party_canonical", "candidate",
    "votes", "candidate_vote_share", "candidate_rank", "candidate_elected",
    "seats_contested", "people_who_voted", "registered_voters",
    "turnout_percent", "turnout_data_source", "turnout_is_reliable",
    "event_type", "polling_date", "source_url", "source_table_index",
    "source_caption", "source_heading", "source_section",
    "total_candidate_votes", "ward_index_status",
]


def load_party_lookup():
    lookup = {}
    for row in csv.DictReader(PARTY_LOOKUP_PATH.open()):
        lookup[row["Party Name As Published"]] = row["Standardised Party Name"]
    return lookup


def parse_ward_sheet(ws):
    """Identical parsing logic to convert_2026_extractor_output.py's
    parse_ward_sheet - locates the Index metadata block, Table 1
    (candidates) and Table 2 (voting summary) by their header rows
    rather than fixed row numbers."""
    rows = list(ws.iter_rows(values_only=True))

    status = None
    for r in rows:
        if r and r[0] == "Extraction status":
            status = r[1]
            break

    candidates, summary = [], {}
    section = None
    for r in rows:
        first = r[0] if r else None
        if first == "Candidate" and r[1] == "Party":
            section = "candidates"
            continue
        if first == "Detail" and r[1] == "Number":
            section = "summary"
            continue
        if first in (None, "") and all(v is None for v in r):
            section = None
            continue
        if section == "candidates" and first:
            candidates.append({
                "candidate": first, "party": r[1], "votes": r[2],
                "vote_share": r[3], "outcome": r[4],
            })
        elif section == "summary" and first:
            summary[first] = r[1]

    return status, candidates, summary


def convert(party_lookup, log_rows):
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True)
    idx = wb["Index"]
    index_rows = [
        {"ward": r[0], "worksheet": r[1], "source_url": r[4]}
        for r in idx.iter_rows(min_row=2, values_only=True) if r[0]
    ]

    out_rows = []
    for entry in index_rows:
        ws = wb[entry["worksheet"]]
        status, candidates, summary = parse_ward_sheet(ws)
        if not candidates:
            log_rows.append({"ward": entry["ward"],
                             "issue": "no_candidates_parsed",
                             "detail": f"worksheet={entry['worksheet']}"})
            continue

        ranked = sorted(candidates, key=lambda c: (-int(c["votes"]), c["candidate"]))
        rank_of = {id(c): i + 1 for i, c in enumerate(ranked)}
        total_candidate_votes = sum(int(c["votes"]) for c in candidates)

        seats = summary.get("Seats")
        electorate = summary.get("Electorate")
        # Genuinely absent from the 2013 official pages (see module
        # docstring) - left blank, not computed from anything else.
        ballot_papers_issued = summary.get("Ballot papers issued")
        turnout = summary.get("Turnout")

        for c in candidates:
            raw_party = c["party"]
            if raw_party is None:
                canonical = None
                log_rows.append({"ward": entry["ward"],
                                 "issue": "candidate_party_not_published",
                                 "detail": c["candidate"]})
            else:
                canonical = party_lookup.get(raw_party)
                if canonical is None:
                    canonical = raw_party
                    log_rows.append({"ward": entry["ward"],
                                     "issue": "party_not_in_standardisation_table",
                                     "detail": raw_party})
            out_rows.append({
                "year": "2013", "council": COUNCIL,
                "ward": entry["ward"], "party_raw": raw_party or "",
                "party_canonical": canonical or "", "candidate": c["candidate"],
                "votes": c["votes"], "candidate_vote_share": c["vote_share"],
                "candidate_rank": rank_of[id(c)],
                "candidate_elected": c["outcome"] == "Elected",
                "seats_contested": seats,
                "people_who_voted": ballot_papers_issued or "",
                "registered_voters": electorate,
                "turnout_percent": turnout or "",
                "turnout_data_source": "" if not turnout else
                                      "official_result_page_via_serpapi_extractor",
                "turnout_is_reliable": bool(turnout),
                "event_type": "scheduled", "polling_date": POLLING_DATE,
                "source_url": entry["source_url"],
                "source_table_index": "", "source_caption": "",
                "source_heading": "", "source_section": "",
                "total_candidate_votes": total_candidate_votes,
                "ward_index_status": status,
            })

        if status != "Complete":
            log_rows.append({
                "ward": entry["ward"], "issue": "index_status_not_complete",
                "detail": f"status={status} (candidate/vote fields present "
                          "and used regardless; turnout/ballot_papers_issued "
                          "genuinely absent from source, not inferred)",
            })

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=OUTPUT_COLUMNS)
        w.writeheader()
        w.writerows(out_rows)
    print(f"{len(out_rows)} candidate rows -> {OUT}")


def main():
    party_lookup = load_party_lookup()
    log_rows = []
    convert(party_lookup, log_rows)

    LOG_OUT.parent.mkdir(parents=True, exist_ok=True)
    with LOG_OUT.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["ward", "issue", "detail"])
        w.writeheader()
        w.writerows(log_rows)

    by_issue = {}
    for r in log_rows:
        by_issue[r["issue"]] = by_issue.get(r["issue"], 0) + 1
    print(f"{len(log_rows)} conversion notes -> {LOG_OUT}")
    print("by issue type:", by_issue)
    print("NOTE: turnout_percent/people_who_voted left blank for every "
          "2013 row - genuinely absent from the official pages. A "
          "follow-up (already approved: Wikipedia turnout, cross-"
          "validated against >=10 official examples) can join those in "
          "later without touching this conversion.")


if __name__ == "__main__":
    main()
