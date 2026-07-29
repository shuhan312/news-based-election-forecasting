"""Export the three news-window schemes side by side, as a workbook.

Three documents specify three different sets of pre-election windows and they
do not reconcile. Arguing from the definitions is slow; showing the same
article landing in three different places is not. This writes that comparison
out so it can be read, sent and decided on.

Four sheets:

    Scheme Definitions   what each scheme says, and where it came from
    Day By Day           every day from polling day back to 181, under all three
    Disagreements        only the days where the schemes differ
    Corpus Impact        the collected articles, counted under each scheme

The last sheet is the one that matters for the decision. Two of the three
schemes exclude everything published more than 30 days before polling, and the
corpus was collected on a 180-day window, so the choice decides how much of
the collection can enter a model at all.

    .venv/bin/python scripts/export_window_scheme_comparison.py
"""

from __future__ import annotations

import json
import sys
from collections import Counter
from datetime import date, time
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from news_modelling.window_schemes import (  # noqa: E402
    SCHEMES,
    assign,
    scheme_summary,
)

OUT_DIRECTORY = REPO / "outputs/window_scheme_comparison"

# A polling day to work the illustration against. 6 May 2021 is a real one and
# has news either side of it, so the day-by-day table is not hypothetical.
REFERENCE_POLL = date(2021, 5, 6)

SCHEME_ORDER = ["original_email_180d", "prompt_1_and_2_30d", "desktop_spec_30d"]
SCHEME_LABELS = {
    "original_email_180d": "Six windows to 180 days",
    "prompt_1_and_2_30d": "Three windows to 30 days",
    "desktop_spec_30d": "Seven periods to 30 days",
}

HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
DISAGREE_FILL = PatternFill("solid", fgColor="FFF2CC")


def day_by_day() -> list[dict]:
    """Every day from polling day back to 181, under all three schemes."""

    rows = []
    for days in range(0, 182):
        published = date.fromordinal(REFERENCE_POLL.toordinal() - days)
        # A time is supplied for polling day so the scheme that has an
        # election-day window can place it; the others exclude it either way.
        published_time = time(9, 0) if days == 0 else None
        row: dict[str, object] = {
            "days before polling": days,
            "publication date": published.isoformat(),
        }
        placements = []
        for key in SCHEME_ORDER:
            result = assign(published, REFERENCE_POLL, scheme=key,
                            published_time=published_time)
            label = result.window or f"excluded: {result.exclusion_reason}"
            row[SCHEME_LABELS[key]] = label
            placements.append(label)
        row["schemes agree"] = "yes" if len(set(placements)) == 1 else "no"
        rows.append(row)
    return rows


def corpus_impact() -> list[dict]:
    """How the collected corpus falls under each scheme.

    Counts real article dates against real polling days rather than the
    illustrative one, so the numbers describe this project's corpus. Articles
    whose publication date could not be established are counted separately -
    they are excluded under every scheme, and lumping them in with articles
    excluded for being too old would misattribute the loss.
    """

    records = REPO / "data/raw/news/records"
    polling = {
        "SCC-2013-05": date(2013, 5, 2), "SCC-2017-05": date(2017, 5, 4),
        "SCC-2021-05": date(2021, 5, 6), "ESWS-2026-05": date(2026, 5, 7),
    }

    counts: dict[str, Counter] = {key: Counter() for key in SCHEME_ORDER}
    undated = 0
    total = 0

    for path in records.glob("*.json"):
        try:
            record = json.loads(path.read_text(encoding="utf-8"))
        except Exception:  # noqa: BLE001 - a malformed record is not a scheme question
            continue
        election = str(record.get("discovered_for_election", ""))
        # Only the four principal elections have a fixed polling day in this
        # table; by-election articles are counted once collection completes.
        poll = next((v for k, v in polling.items()
                     if k.split("-")[1] in election and "by-election" not in election),
                    None)
        if poll is None:
            continue
        total += 1

        raw = record.get("dates", {}).get("published_date")
        if not raw:
            undated += 1
            continue
        try:
            published = date.fromisoformat(str(raw)[:10])
        except ValueError:
            undated += 1
            continue

        for key in SCHEME_ORDER:
            result = assign(published, poll, scheme=key)
            counts[key]["included" if result.included else "excluded"] += 1

    rows = [{
        "scheme": SCHEME_LABELS[key],
        "articles included in influence features": counts[key]["included"],
        "articles excluded by timing": counts[key]["excluded"],
        "share of dated articles usable":
            round(counts[key]["included"] / max(sum(counts[key].values()), 1), 3),
    } for key in SCHEME_ORDER]

    rows.append({
        "scheme": "— articles with no usable publication date —",
        "articles included in influence features": 0,
        "articles excluded by timing": undated,
        "share of dated articles usable": "excluded under every scheme",
    })
    rows.append({
        "scheme": "— principal-election articles considered —",
        "articles included in influence features": total,
        "articles excluded by timing": "",
        "share of dated articles usable": "",
    })
    return rows


def write(workbook: Workbook, title: str, rows: list[dict],
          highlight_column: str | None = None) -> None:
    sheet = workbook.create_sheet(title=title[:31])
    if not rows:
        sheet.append(["(no rows)"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row in rows:
        sheet.append([
            ", ".join(map(str, value)) if isinstance(value, list) else value
            for value in (row.get(header) for header in headers)
        ])
        # Shade the rows where the schemes disagree, so the pattern is visible
        # without reading every cell. The status is also a text column, so the
        # shading is a convenience rather than the only signal.
        if highlight_column and str(row.get(highlight_column)) == "no":
            for cell in sheet[sheet.max_row]:
                cell.fill = DISAGREE_FILL

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for index, header in enumerate(headers, start=1):
        widest = max([len(str(header))]
                     + [len(str(row.get(header, ""))) for row in rows])
        sheet.column_dimensions[get_column_letter(index)].width = min(widest + 2, 48)


def main() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    write(workbook, "Scheme Definitions", [
        {
            "scheme": SCHEME_LABELS[row["scheme"]],
            "key": row["scheme"],
            "stated in": row["source"],
            "non-overlapping windows": row["windows"],
            "cumulative snapshots": row["cumulative"],
            "modelled horizon (days)": row["modelled_horizon_days"],
            "has an election-day window": row["includes_election_day"],
        }
        for row in scheme_summary()
    ])

    days = day_by_day()
    write(workbook, "Day By Day", days, highlight_column="schemes agree")
    write(workbook, "Disagreements",
          [row for row in days if row["schemes agree"] == "no"],
          highlight_column="schemes agree")
    write(workbook, "Corpus Impact", corpus_impact())

    OUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    path = OUT_DIRECTORY / "News_Window_Scheme_Comparison.xlsx"
    workbook.save(path)

    disagreements = sum(1 for row in days if row["schemes agree"] == "no")
    print(f"written to {path}")
    print(f"  {len(days)} days compared, {disagreements} where the schemes differ")
    for row in corpus_impact()[:3]:
        print(f"  {row['scheme']:28s} "
              f"{row['articles included in influence features']:6d} usable, "
              f"{row['articles excluded by timing']:6d} excluded by timing")


if __name__ == "__main__":
    main()
